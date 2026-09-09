# -*- coding: utf-8 -*-
"""P1 PPT 入库验收：
- .pptx 上传入库：文字页零费用提取 + 图片型页 OCR 兜底
- chunk 携带页码、检索命中图片 OCR 文本、HTTP 上传端点接受 pptx
用法：在 app 目录下运行  venv\\Scripts\\python.exe tests\\test_ppt.py
（需 API Key：embedding 少量 + OCR 约 1~2 分/图）
"""
import os
import sys
import tempfile
import time

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import config  # noqa: E402
from backend import kb as kb_mod  # noqa: E402
from backend.server import create_app, load_settings  # noqa: E402

FAILED = []


def check(name, ok, detail=""):
    print(("PASS | " if ok else "FAIL | ") + name + (f" | {detail}" if detail else ""))
    if not ok:
        FAILED.append(name)


def wait_ingest(kb_id, timeout=300):
    deadline = time.time() + timeout
    while time.time() < deadline:
        st = kb_mod.get_ingest_status(kb_id) or {}
        if not st.get("running"):
            return st
        time.sleep(0.5)
    return {"error": "ingest timeout"}


def _load_font(size):
    """找系统中文字体渲染测试图（PIL 默认位图字体不支持中文，会画成方块）。"""
    from PIL import ImageFont
    for fp in (r"C:\Windows\Fonts\msyh.ttc", r"C:\Windows\Fonts\simhei.ttf", r"C:\Windows\Fonts\simsun.ttc",
               "/System/Library/Fonts/PingFang.ttc", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"):
        if os.path.exists(fp):
            try:
                return ImageFont.truetype(fp, size)
            except Exception:
                pass
    return ImageFont.load_default()


def make_test_pptx(path):
    """生成测试课件：第1页文字（含备注）、第2页纯图片（需 OCR 兜底）、第3页文字+图片（不触发 OCR）。"""
    from pptx import Presentation
    from pptx.util import Inches
    from PIL import Image, ImageDraw

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "PPT入库验收：文字页"
    slide.notes_slide.notes_text_frame.text = "备注要点：文字页应零费用直接入库"

    img_path = path + "_img.png"
    img = Image.new("RGB", (640, 240), "white")
    d = ImageDraw.Draw(img)
    font = _load_font(28)
    d.text((40, 60), "PPT图片页OCR验收文字：道尔顿分压定律 P总 = P1 + P2", fill="black", font=font)
    d.text((40, 120), "第二行：体积功 W = -P外dV", fill="black", font=font)
    img.save(img_path)

    slide = prs.slides.add_slide(prs.slide_layouts[6])  # 空白布局：无任何文本层
    slide.shapes.add_picture(img_path, Inches(1), Inches(1), width=Inches(8))

    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "PPT入库验收：混合页标题文字"
    slide.shapes.add_picture(img_path, Inches(1), Inches(2), width=Inches(8))

    prs.save(path)
    os.unlink(img_path)


def main():
    kb_mod.ensure_data_dirs()
    settings = load_settings()
    if not settings.get("api_key"):
        print("!! 尚未配置 API Key")
        sys.exit(1)
    app = create_app()
    client = app.test_client()

    tmp = tempfile.NamedTemporaryFile("wb", suffix=".pptx", delete=False)
    tmp.close()
    make_test_pptx(tmp.name)

    kb_id = kb_mod.create_kb("PPT测试")
    try:
        kb_mod.import_files(kb_id, [tmp.name])
        kb_mod.start_ingest(kb_id, settings)
        st = wait_ingest(kb_id)
        check("PPT 索引完成无报错", not st.get("error") and not (st.get("stats") or {}).get("errors"), str(st)[:300])
        stats = st.get("stats") or {}
        if st.get("error"):
            print("   索引失败，退出")
            sys.exit(1)
        check("chunks >= 3（每页至少一块）", stats.get("chunks", 0) >= 3, str(stats))
        check("图片页 OCR 无跳过", not stats.get("ocr_skipped"), str(stats.get("ocr_skipped")))

        # 检索：图片页 OCR 文本应命中
        r = client.post("/api/search", json={
            "question": "道尔顿分压定律", "kb_ids": [kb_id], "top_k": 6, "threshold": 0.0,
        }).get_json()
        hits = r.get("hits") or []
        top = hits[0] if hits else {}
        check("检索命中图片页 OCR 文本",
              bool(hits) and any(("道尔顿" in h["text"] or "分压" in h["text"]) for h in hits),
              str(top.get("text", ""))[:100])
        check("命中块带页码", bool(top.get("meta", {}).get("page")), str(top.get("meta", {})))

        # 检索：文字页内容（含备注）
        r = client.post("/api/search", json={
            "question": "PPT入库验收 文字页", "kb_ids": [kb_id], "top_k": 6, "threshold": 0.0,
        }).get_json()
        check("检索命中文字页内容", any("文字页" in h["text"] for h in (r.get("hits") or [])), "")

        # 文档列表显示 .pptx
        docs = client.get(f"/api/kbs/{kb_id}/docs").get_json()["docs"]
        check("文档列表包含 pptx", any(d["rel"].endswith(".pptx") for d in docs), str(docs))

        # HTTP 上传端点冒烟（accept .pptx 白名单）
        kb2 = kb_mod.create_kb("PPT上传测试")
        try:
            with open(tmp.name, "rb") as f:
                r = client.post(f"/api/kbs/{kb2}/upload", data={"files": (f, "上传课件.pptx")},
                                content_type="multipart/form-data")
            check("HTTP 上传 pptx 成功",
                  r.get_json().get("ok") is True and r.get_json().get("imported") == 1,
                  str(r.get_json()))
            wait_ingest(kb2)  # 等上传触发的索引结束再删除，避免后台线程写已删目录
        finally:
            kb_mod.delete_kb(kb2)
    finally:
        kb_mod.delete_kb(kb_id)
        try:
            os.unlink(tmp.name)
        except OSError:
            pass

    print("\n===== P1 PPT result: %s =====" % ("ALL PASS" if not FAILED else f"{len(FAILED)} FAILED: {FAILED}"))
    sys.exit(1 if FAILED else 0)


if __name__ == "__main__":
    main()
