# -*- coding: utf-8 -*-
"""笔记导出：把当前对话（含引用原文）导出为 Markdown 或 Word 文件。"""
import os
import time


def _safe_name(title):
    bad = '<>:"/\\|?*'
    name = "".join("_" if c in bad else c for c in (title or "对话笔记")).strip() or "对话笔记"
    return name[:60]


def _default_dir():
    for candidate in (os.path.expanduser("~/Downloads"), os.path.expanduser("~/Documents"), os.path.expanduser("~")):
        if os.path.isdir(candidate):
            return candidate
    return os.path.expanduser("~")


def _build_sections(messages):
    """把消息整理成 [{role, content, citations}]，便于两种格式共用。"""
    out = []
    for m in messages:
        if m["role"] not in ("user", "assistant"):
            continue
        out.append({"role": m["role"], "content": m["content"], "citations": m.get("citations") or []})
    return out


def export_conversation(conv, messages, fmt="md", out_path=None):
    title = conv.get("title") or "对话笔记"
    if fmt == "docx":
        ext = ".docx"
    else:
        ext = ".md"
    if out_path is None:
        out_path = os.path.join(_default_dir(), _safe_name(title) + ext)
    sections = _build_sections(messages)
    if fmt == "docx":
        _write_docx(title, sections, out_path)
    else:
        _write_markdown(title, sections, out_path)
    return out_path


def export_note_text(title, content, fmt="md", out_path=None):
    """单篇知识笔记导出（Markdown / Word），复用对话导出的写入逻辑。"""
    ext = ".docx" if fmt == "docx" else ".md"
    if out_path is None:
        out_path = os.path.join(_default_dir(), _safe_name(title) + ext)
    sections = [{"role": "note", "content": content, "citations": []}]
    if fmt == "docx":
        _write_docx(title, sections, out_path)
    else:
        _write_markdown(title, sections, out_path)
    return out_path


# ---------- Markdown ----------

def _write_markdown(title, sections, out_path):
    lines = [f"# {title}", "", f"> 导出时间：{time.strftime('%Y-%m-%d %H:%M')}｜由「实验室助手」生成", ""]
    for i, s in enumerate(sections, 1):
        if s["role"] == "user":
            lines.append(f"## 问题 {i}：{s['content']}")
        elif s["role"] == "note":
            lines.append(s["content"])  # 知识笔记：正文本身就是完整 Markdown 文档
        else:
            lines.append(f"### 回答 {i}")
            lines.append("")
            lines.append(s["content"])
            if s["citations"]:
                lines.append("")
                lines.append("**引用来源（原文）**")
                for c in s["citations"]:
                    src = "｜".join(x for x in (c.get("kb_name", ""), c.get("category", ""), c.get("source", "")) if x)
                    lines.append(f"\n> [{c['n']}]（{src}，相似度 {c['score']}）")
                    lines.append(f"> {c['text']}")
            lines.append("")
    with open(out_path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))


# ---------- Word ----------

def _write_docx(title, sections, out_path):
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.oxml.ns import qn

    doc = Document()
    doc.add_heading(title, level=0)
    p = doc.add_paragraph(f"导出时间：{time.strftime('%Y-%m-%d %H:%M')}｜由「实验室助手」生成")
    p.runs[0].font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    for i, s in enumerate(sections, 1):
        if s["role"] == "user":
            doc.add_heading(f"问题 {i}：{s['content']}", level=1)
        elif s["role"] == "note":
            # 知识笔记：按行写入，以 # 开头的行转标题层级
            for line in s["content"].splitlines():
                if line.startswith("### "):
                    doc.add_heading(line[4:], level=3)
                elif line.startswith("## "):
                    doc.add_heading(line[3:], level=2)
                elif line.startswith("# "):
                    doc.add_heading(line[2:], level=1)
                else:
                    doc.add_paragraph(line)
        else:
            doc.add_heading(f"回答 {i}", level=2)
            doc.add_paragraph(s["content"])
            if s["citations"]:
                doc.add_heading("引用来源（原文）", level=3)
                for c in s["citations"]:
                    src = "｜".join(x for x in (c.get("kb_name", ""), c.get("category", ""), c.get("source", "")) if x)
                    q = doc.add_paragraph()
                    q.paragraph_format.left_indent = Pt(18)
                    run = q.add_run(f"[{c['n']}]（{src}，相似度 {c['score']}）\n{c['text']}")
                    run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
                    run.font.size = Pt(10)
                    # 中文用楷体，西文用 Consolas，风格接近"引用块"
                    run.font.name = "Consolas"
                    run._element.rPr.rFonts.set(qn("w:eastAsia"), "楷体")
    doc.save(out_path)
