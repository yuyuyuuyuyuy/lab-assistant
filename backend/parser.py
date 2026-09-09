# -*- coding: utf-8 -*-
"""文档解析：txt / pdf / docx → 分块列表。

分块规则（Dify 项目验证过的结论）：
- txt 按行切分（语料按"每行一条完整知识点"编写，每行自带主题词+数值）；
- pdf/docx 按自然段落切分，短段合并、超长段按句号分号断句；
- 绝不按固定字符数硬切。
"""
import os

from charset_normalizer import from_bytes


def read_text_file(path):
    """读取 txt，自动探测编码（UTF-8 / GB18030 等），失败时 ignore 兜底防乱码崩溃。"""
    with open(path, "rb") as f:
        raw = f.read()
    if not raw:
        return ""
    best = from_bytes(raw).best()
    if best is None:
        return raw.decode("utf-8", errors="ignore")
    return str(best)


def split_paragraphs(text, max_len=800, min_len=8):
    """按自然段落切块：短段落合并，超长段落按句号/分号断句。"""
    paras = []
    for p in text.replace("\r\n", "\n").split("\n\n"):
        joined = " ".join(l.strip() for l in p.split("\n") if l.strip())
        if joined:
            paras.append(joined)

    chunks, buf = [], ""
    for p in paras:
        buf = (buf + "\n" + p).strip()
        if len(buf) > max_len:
            chunks.append(buf)
            buf = ""
    if buf:
        chunks.append(buf)

    result = []
    for c in chunks:
        if len(c) <= max_len:
            result.append(c)
            continue
        pieces, cur = [], ""
        for ch in c:
            cur += ch
            if ch in "。；;" and len(cur) >= min_len:
                pieces.append(cur.strip())
                cur = ""
        if cur.strip():
            pieces.append(cur.strip())
        result.extend(pieces)
    return [r for r in result if len(r) >= min_len]


def parse_txt(path):
    """txt 按行切分，每行一条完整知识点。"""
    text = read_text_file(path)
    return [l.strip() for l in text.splitlines() if len(l.strip()) >= 4]


def parse_pdf(path):
    import pymupdf

    doc = pymupdf.open(path)
    try:
        pages_text, scanned = [], 0
        for page in doc:
            t = page.get_text().strip()
            pages_text.append(t)
            if len(t) < 30:
                scanned += 1
    finally:
        doc.close()
    if pages_text and scanned / len(pages_text) > 0.7:
        raise ValueError("疑似扫描版 PDF（缺少文字层），请在知识库中使用「整本 OCR」识别后入库")
    return split_paragraphs("\n\n".join(pages_text))


OCR_SIDECAR_EXT = ".ocr.json"  # 整本 OCR 结果：与原 PDF 同目录的「文件名.pdf.ocr.json」


def parse_ocr_json(path):
    """解析整本 OCR 的 sidecar 文件，返回 [(页面文本, 页码)]（空页跳过）。"""
    import json

    with open(path, encoding="utf-8") as f:
        data = json.load(f)
    out = []
    for p in data.get("pages") or []:
        text = (p.get("text") or "").strip()
        if text:
            out.append((text, p.get("page")))
    return out


def parse_docx(path):
    from docx import Document

    paras = [p.text.strip() for p in Document(path).paragraphs if p.text.strip()]
    if not paras:
        return []
    return split_paragraphs("\n\n".join(paras))


def parse_pptx(path, image_min_chars=10):
    """解析 PPT（.pptx）：提取每页文本层（正文/表格/备注/图表数据），
    返回 [(text, page, images)]，page 从 1 开始。

    图片型页（文本不足 image_min_chars）附上该页内嵌图片 blob，供调用方
    送 OCR 兜底；文本足够的页 images 为空列表（零 API 费用）。
    """
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    def shape_texts(shape):
        """递归提取单个形状的文本：组 / 表格 / 图表 / 文本框。"""
        texts = []
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            for sub in shape.shapes:
                texts.extend(shape_texts(sub))
            return texts
        if getattr(shape, "has_table", False) and shape.has_table:
            for row in shape.table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    texts.append("｜".join(cells))
            return texts
        if getattr(shape, "has_chart", False) and shape.has_chart:
            try:
                bits = []
                for plot in shape.chart.plots:
                    try:
                        bits.append("、".join(str(c) for c in plot.categories))
                    except Exception:
                        pass
                    for series in plot.series:
                        try:
                            bits.append("、".join(str(v) for v in series.values))
                        except Exception:
                            pass
                if bits:
                    texts.append("图表数据：" + "；".join(bits))
            except Exception:
                pass
        if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
            t = shape.text_frame.text.strip()
            if t:
                texts.append(t)
        return texts

    def collect_images(shape):
        """递归收集形状内嵌图片的原始字节（含组内图片）。"""
        imgs = []
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            for sub in shape.shapes:
                imgs.extend(collect_images(sub))
        elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            try:
                imgs.append(shape.image.blob)
            except Exception:
                pass
        return imgs

    prs = Presentation(path)
    out = []
    for idx, slide in enumerate(prs.slides, 1):
        texts = []
        for shape in slide.shapes:
            texts.extend(shape_texts(shape))
        if slide.has_notes_slide:  # 备注往往是老师讲稿要点，并入页尾
            note = slide.notes_slide.notes_text_frame.text.strip()
            if note:
                texts.append(note)
        text = "\n".join(texts).strip()
        images = []
        if len(text) < image_min_chars:
            for shape in slide.shapes:
                images.extend(collect_images(shape))
        out.append((text, idx, images))
    return out


SUPPORTED_EXTS = {".txt": parse_txt, ".pdf": parse_pdf, ".docx": parse_docx, ".pptx": parse_pptx}


def parse_file(path):
    ext = os.path.splitext(path)[1].lower()
    if ext not in SUPPORTED_EXTS:
        raise ValueError(f"不支持的文件格式 {ext}（仅支持 txt / pdf / docx / pptx）")
    if ext == ".pptx":
        # 泛化路径的兜底：丢弃页码与图片（图片页 OCR 由 ingest 的 pptx 专用分支处理）
        return [t for t, _p, _i in parse_pptx(path) if t]
    return SUPPORTED_EXTS[ext](path)
